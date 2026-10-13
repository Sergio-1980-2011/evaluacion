"""Genera la planilla de evaluación de la rotación de prehabilitación.

Uso: python3 generar_planilla.py  (escribe ../Planilla_Evaluacion_Rotacion.xlsx)
"""
import datetime
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
DATOS = json.loads((AQUI / "contenido.json").read_text(encoding="utf-8"))
SALIDA = AQUI.parent / "Planilla_Evaluacion_Rotacion.xlsx"

FUENTE = "Arial"
AZUL_OSCURO = "1F4E79"
AZUL_CLARO = "DDEBF7"
AMARILLO = "FFF2CC"
GRIS = "F2F2F2"
VERDE = "C6EFCE"
ROJO = "FFC7CE"

fino = Side(style="thin", color="BFBFBF")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)


def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FUENTE, bold=bold, size=size, color=color, italic=italic)


def relleno(color):
    return PatternFill("solid", start_color=color, end_color=color)


def titulo(ws, texto, sub=None, ancho_cols=8):
    ws["A1"] = texto
    ws["A1"].font = f(True, 14, AZUL_OSCURO)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(False, 9, "595959", italic=True)
    ws.sheet_view.showGridLines = False


def encabezado(ws, fila, valores, col_ini=1):
    for i, v in enumerate(valores):
        c = ws.cell(row=fila, column=col_ini + i, value=v)
        c.font = f(True, 10, "FFFFFF")
        c.fill = relleno(AZUL_OSCURO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE


def celda(ws, ref, valor=None, entrada=False, bold=False, fmt=None, centro=False, wrap=False):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.font = f(bold, 10, "0000FF" if entrada else "000000")
    c.border = BORDE
    if entrada:
        c.fill = relleno(AMARILLO)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(
        horizontal="center" if centro else None, vertical="center", wrap_text=wrap
    )
    return c


wb = Workbook()

# ---------------------------------------------------------------- Parámetros
par = wb.active
par.title = "Parametros"
titulo(par, "Parámetros de la evaluación", "Valores definidos por el programa. Modificar sólo si cambia el reglamento.")
encabezado(par, 4, ["Parámetro", "Valor", "Fuente / nota"])
parametros = [
    ("Nota mínima de la escala", 1, "0.0", "Escala de notas 1 a 7 (programa de la rotación)."),
    ("Nota máxima de la escala", 7, "0.0", "Escala de notas 1 a 7 (programa de la rotación)."),
    ("Nota de aprobación de la escala (nota 4,0)", 4, "0.0", "Nota que se obtiene con el porcentaje de exigencia."),
    ("Porcentaje de exigencia", DATOS["escala"]["exigencia"], "0%", "Supuesto: 60 % de logro = nota 4,0 (estándar habitual en Chile). Ajustar si el reglamento indica otro."),
    ("Puntaje máximo por criterio", 4, "0", "Niveles de desempeño 0 a 4 (ver hojas de rúbricas)."),
    ("Ponderación Seminarios", 0.5, "0%", "Programa: Pauta de presentación oral = 50 % de la nota final."),
    ("Ponderación Desempeño (Pauta global)", 0.5, "0%", "Programa: Pauta de evaluación global = 50 % de la nota final."),
    ("Nota final mínima de aprobación", 5.0, "0.0", "Programa: requisito 2."),
    ("Nota mínima en Seminarios", 5.0, "0.0", "Programa: requisito 3."),
    ("Asistencia mínima", 0.9, "0%", "Programa: requisito 1."),
]
for i, (nombre, valor, fmt, nota) in enumerate(parametros, start=5):
    celda(par, f"A{i}", nombre)
    celda(par, f"B{i}", valor, entrada=True, fmt=fmt, centro=True)
    celda(par, f"C{i}", nota, wrap=True)
par.column_dimensions["A"].width = 42
par.column_dimensions["B"].width = 12
par.column_dimensions["C"].width = 80
P_MIN, P_MAX, P_APR, P_EXI, P_PMAX, P_WSEM, P_WDES, P_NFIN, P_NSEM, P_ASIS = (
    f"Parametros!$B${i}" for i in range(5, 15)
)


def formula_nota(pct):
    """Nota chilena con exigencia: tramo lineal bajo y sobre la exigencia, redondeada a 1 decimal."""
    return (
        f"ROUND(IF({pct}<{P_EXI},{P_MIN}+({P_APR}-{P_MIN})*{pct}/{P_EXI},"
        f"{P_APR}+({P_MAX}-{P_APR})*({pct}-{P_EXI})/(1-{P_EXI})),1)"
    )


# ---------------------------------------------------------------- Instrucciones
ins = wb.create_sheet("Instrucciones", 0)
titulo(ins, "Planilla de evaluación – Rotación de Prehabilitación Quirúrgica",
       "Residente de Anestesiología · Servicio de Kinesiología · Noviembre 2026")
ins.column_dimensions["A"].width = 4
ins.column_dimensions["B"].width = 30
ins.column_dimensions["C"].width = 95
filas = [
    ("Cómo usar", None),
    ("1. Resumen", "Completar los datos del residente y docentes. La nota final y el estado de aprobación se calculan solos."),
    ("2. Asistencia", "Marcar cada día: P = presente, A = ausente, J = ausencia justificada, R = día recuperado. Asistencia = (P + R) / días programados."),
    ("3. Seminarios", "Puntuar de 0 a 4 cada criterio de los 4 seminarios (descriptores en la hoja «Rubrica_Seminario»). La nota aparece cuando los 8 criterios tienen puntaje."),
    ("4. Desempeño", "Puntuar de 0 a 4 cada criterio (descriptores en «Rubrica_Desempeno»). La columna «Mitad de rotación» es formativa y NO entra en la nota; sólo cuenta la columna «Final»."),
    ("5. Bitácora", "Registrar la entrega semanal en Canvas y un comentario breve del docente."),
    ("6. Feedback", "Registrar cada sesión de feedback: autoevaluación, fortalezas, aspectos a mejorar y plan de acción."),
    ("7. Escala_Notas", "Tabla de conversión % de logro → nota, calculada con el % de exigencia de «Parametros»."),
    ("Código de colores", None),
    ("Celda amarilla, texto azul", "Dato que se debe ingresar."),
    ("Texto negro, sin relleno", "Fórmula: no modificar."),
    ("Verde / rojo", "Formato condicional: cumple / no cumple un requisito de aprobación."),
    ("Cálculo de la nota", None),
    ("% de logro", "Σ (ponderación del criterio × puntaje / 4). Los pesos de cada criterio suman 100 %."),
    ("Nota", "Si % < exigencia: nota = 1 + 3 × %/exigencia. Si % ≥ exigencia: nota = 4 + 3 × (% − exigencia)/(1 − exigencia). Se redondea a 1 decimal."),
    ("Nota final", "50 % promedio de seminarios + 50 % evaluación global final (redondeada a 1 decimal)."),
    ("Aprobación", "Asistencia ≥ 90 %  ·  Nota final ≥ 5,0  ·  Nota de seminarios ≥ 5,0. Los tres requisitos son obligatorios."),
]
r = 4
for a, b in filas:
    if b is None:
        r += 1
        ins.cell(row=r, column=2, value=a).font = f(True, 11, AZUL_OSCURO)
    else:
        ins.cell(row=r, column=2, value=a).font = f(True)
        if a.startswith("Celda amarilla"):
            ins.cell(row=r, column=2).fill = relleno(AMARILLO)
            ins.cell(row=r, column=2).font = f(True, color="0000FF")
        c = ins.cell(row=r, column=3, value=b)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
r += 1
ins.cell(row=r, column=2, value="Ejemplo de cálculo").font = f(True, 11, AZUL_OSCURO)
r += 1
encabezado(ins, r, ["Criterio (pauta seminario)", "Puntaje de ejemplo (0–4)", "Peso"], col_ini=2)
ej_ini = r + 1
ejemplo = [4, 3, 3, 2, 4, 3, 3, 3]
for i, (crit, pts) in enumerate(zip(DATOS["pauta_seminario"], ejemplo)):
    rr = ej_ini + i
    celda(ins, f"B{rr}", crit["criterio"], wrap=True)
    celda(ins, f"C{rr}", pts, entrada=True)
    ins[f"D{rr}"] = crit["peso"]
    ins[f"D{rr}"].number_format = "0%"
    ins[f"D{rr}"].font = f(color="808080")
ej_fin = ej_ini + len(ejemplo) - 1
ins.cell(row=ej_fin + 1, column=2, value="% de logro").font = f(True)
celda(ins, f"C{ej_fin + 1}", f"=SUMPRODUCT(D{ej_ini}:D{ej_fin},C{ej_ini}:C{ej_fin})/{P_PMAX}", fmt="0.0%")
ins.cell(row=ej_fin + 2, column=2, value="Nota").font = f(True)
celda(ins, f"C{ej_fin + 2}", "=" + formula_nota(f"C{ej_fin + 1}"), fmt="0.0", bold=True)
ins.column_dimensions["D"].width = 8

# ---------------------------------------------------------------- Resumen
res = wb.create_sheet("Resumen", 1)
titulo(res, "Resumen y nota final de la rotación")
res.column_dimensions["A"].width = 44
res.column_dimensions["B"].width = 34
res.column_dimensions["C"].width = 16
datos_id = [
    ("Residente", "Nombre Apellido"),
    ("RUT", ""),
    ("Programa", "Anestesiología y Reanimación"),
    ("Docente guía", ""),
    ("Kinesiólogo(a) supervisor(a)", ""),
    ("Período", "02-11-2026 al 30-11-2026"),
]
encabezado(res, 3, ["Identificación", "Dato"])
for i, (k, v) in enumerate(datos_id, start=4):
    celda(res, f"A{i}", k, bold=True)
    celda(res, f"B{i}", v, entrada=True)
res["B4"].comment = Comment("Ejemplo de formato: reemplazar por el nombre del residente.", "Planilla")

encabezado(res, 11, ["Componente", "Resultado", "Ponderación"])
celda(res, "A12", "Nota Seminarios (promedio de 4)")
celda(res, "B12", "=Seminarios!G20", fmt="0.0", centro=True)
celda(res, "C12", f"={P_WSEM}", fmt="0%", centro=True)
celda(res, "A13", "Nota Desempeño (Pauta global final)")
celda(res, "B13", "=Desempeno!E19", fmt="0.0", centro=True)
celda(res, "C13", f"={P_WDES}", fmt="0%", centro=True)
celda(res, "A14", "NOTA FINAL DE LA ROTACIÓN", bold=True)
celda(res, "B14", '=IF(AND(ISNUMBER(B12),ISNUMBER(B13)),ROUND(B12*C12+B13*C13,1),"")',
      fmt="0.0", centro=True, bold=True)
res["B14"].font = f(True, 14)
celda(res, "A15", "Asistencia")
celda(res, "B15", "=Asistencia!C30", fmt="0.0%", centro=True)

encabezado(res, 17, ["Requisito de aprobación", "Exigido", "¿Cumple?"])
celda(res, "A18", "1. Asistencia mínima")
celda(res, "B18", f"={P_ASIS}", fmt="0%", centro=True)
celda(res, "C18", '=IF(ISNUMBER(B15),IF(B15>=B18,"SÍ","NO"),"PENDIENTE")', centro=True)
celda(res, "A19", "2. Nota final de la rotación")
celda(res, "B19", f"={P_NFIN}", fmt="0.0", centro=True)
celda(res, "C19", '=IF(ISNUMBER(B14),IF(B14>=B19,"SÍ","NO"),"PENDIENTE")', centro=True)
celda(res, "A20", "3. Nota en el área de Seminarios")
celda(res, "B20", f"={P_NSEM}", fmt="0.0", centro=True)
celda(res, "C20", '=IF(ISNUMBER(B12),IF(B12>=B20,"SÍ","NO"),"PENDIENTE")', centro=True)
celda(res, "A22", "ESTADO DE LA ROTACIÓN", bold=True)
celda(res, "B22",
      '=IF(COUNTIF(C18:C20,"NO")>0,"REPROBADA",IF(COUNTIF(C18:C20,"PENDIENTE")>0,"EN CURSO","APROBADA"))',
      centro=True, bold=True)
res["B22"].font = f(True, 12)
for rng in ("C18:C20", "B22"):
    res.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({rng.split(":")[0]}="SÍ",{rng.split(":")[0]}="APROBADA")'], fill=relleno(VERDE)))
    res.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({rng.split(":")[0]}="NO",{rng.split(":")[0]}="REPROBADA")'], fill=relleno(ROJO)))
celda(res, "A24", "Observaciones finales del docente guía", bold=True)
res.merge_cells("A25:C29")
celda(res, "A25", "", entrada=True, wrap=True)
res["A25"].alignment = Alignment(wrap_text=True, vertical="top")
celda(res, "A31", "Firma docente guía:")
celda(res, "A32", "Firma residente:")
celda(res, "A33", "Fecha:")
for rr in (31, 32, 33):
    celda(res, f"B{rr}", "", entrada=True)

# ---------------------------------------------------------------- Asistencia
asi = wb.create_sheet("Asistencia", 2)
titulo(asi, "Registro de asistencia", "P = presente · A = ausente · J = ausencia justificada · R = día recuperado")
encabezado(asi, 4, ["Fecha", "Día", "Estado", "Actividad / observación"])
dias = ["lunes", "martes", "miércoles", "jueves", "viernes"]
fecha = datetime.date(2026, 11, 2)
fila = 5
while fecha <= datetime.date(2026, 11, 30):
    if fecha.weekday() < 5:
        celda(asi, f"A{fila}", fecha, fmt="dd-mm-yyyy", centro=True)
        celda(asi, f"B{fila}", dias[fecha.weekday()], centro=True)
        celda(asi, f"C{fila}", None, entrada=True, centro=True)
        obs = "Seminario" if fecha.weekday() == 4 else ""
        if fecha == datetime.date(2026, 11, 30):
            obs = "Cierre: evaluación final y feedback"
        celda(asi, f"D{fila}", obs, entrada=True)
        fila += 1
    fecha += datetime.timedelta(days=1)
ult = fila - 1  # 25
dv = DataValidation(type="list", formula1='"P,A,J,R"', allow_blank=True)
dv.error = "Usar P, A, J o R"
asi.add_data_validation(dv)
dv.add(f"C5:C{ult}")
celda(asi, "B27", "Días programados", bold=True)
celda(asi, "C27", f"=COUNTA(A5:A{ult})", centro=True)
celda(asi, "B28", "Días asistidos (P + R)", bold=True)
celda(asi, "C28", f'=COUNTIF(C5:C{ult},"P")+COUNTIF(C5:C{ult},"R")', centro=True)
celda(asi, "B29", "Días registrados", bold=True)
celda(asi, "C29", f"=COUNTA(C5:C{ult})", centro=True)
celda(asi, "B30", "% Asistencia", bold=True)
celda(asi, "C30", '=IF(C29=0,"",C28/C27)', fmt="0.0%", centro=True)
celda(asi, "B31", "Máximo de inasistencias permitidas", bold=True)
celda(asi, "C31", f"=C27-ROUNDUP(C27*{P_ASIS},0)", centro=True)
asi["D30"] = "Se calcula sobre el total de días programados; las ausencias justificadas no recuperadas cuentan como inasistencia."
asi["D30"].font = f(size=9, italic=True, color="595959")
asi.conditional_formatting.add(f"C5:C{ult}", CellIsRule(operator="equal", formula=['"A"'], fill=relleno(ROJO)))
asi.conditional_formatting.add(f"C5:C{ult}", CellIsRule(operator="equal", formula=['"J"'], fill=relleno("FFEB9C")))
for col, w in zip("ABCD", (14, 12, 10, 60)):
    asi.column_dimensions[col].width = w


# ---------------------------------------------------------------- Hojas de puntaje
def hoja_puntaje(nombre, tit, pauta, columnas, fila_titulos=None):
    ws = wb.create_sheet(nombre)
    titulo(ws, tit, "Puntaje por criterio: 4 Excelente · 3 Bueno · 2 Suficiente · 1 Insuficiente · 0 No logrado")
    n = len(pauta)
    cab = ["N°", "Criterio", "Peso"] + columnas
    encabezado(ws, 6, cab)
    if fila_titulos:
        for etiqueta, valores, fila in fila_titulos:
            ws.cell(row=fila, column=2, value=etiqueta).font = f(True)
            for j, v in enumerate(valores):
                celda(ws, f"{get_column_letter(4 + j)}{fila}", v, entrada=True, wrap=True, centro=True)
    for i, crit in enumerate(pauta):
        rr = 7 + i
        celda(ws, f"A{rr}", i + 1, centro=True)
        celda(ws, f"B{rr}", crit["criterio"], wrap=True)
        celda(ws, f"C{rr}", crit["peso"], fmt="0%", centro=True)
        for j in range(len(columnas)):
            celda(ws, f"{get_column_letter(4 + j)}{rr}", None, entrada=True, centro=True)
    fin = 6 + n
    dv = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True)
    dv.error = "Ingresar un puntaje entero de 0 a 4"
    ws.add_data_validation(dv)
    dv.add(f"D7:{get_column_letter(3 + len(columnas))}{fin}")
    celda(ws, f"B{fin + 1}", "Suma de ponderaciones", bold=True)
    celda(ws, f"C{fin + 1}", f"=SUM(C7:C{fin})", fmt="0%", centro=True)
    return ws, fin


# Seminarios
sem_cols = [f"Seminario {s['n']}" for s in DATOS["curso"]["seminarios"]]
sem, fin = hoja_puntaje(
    "Seminarios", "Evaluación de seminarios – Pauta de presentación oral",
    DATOS["pauta_seminario"], sem_cols,
    fila_titulos=[("Tema", [s["titulo"] for s in DATOS["curso"]["seminarios"]], 3),
                  ("Fecha", [s["fecha"] for s in DATOS["curso"]["seminarios"]], 4),
                  ("Docente evaluador", ["", "", "", ""], 5)],
)
# fin = 14 -> filas 15 suma, 16 puntaje, 17 % logro, 18 nota, 20 promedio
for j in range(4):
    col = get_column_letter(4 + j)
    rng = f"{col}7:{col}{fin}"
    celda(sem, f"{col}16", f"=SUM({rng})", centro=True)
    celda(sem, f"{col}17", f'=IF(COUNT({rng})<{len(DATOS["pauta_seminario"])},"",SUMPRODUCT($C$7:$C${fin},{rng})/{P_PMAX})',
          fmt="0.0%", centro=True)
    celda(sem, f"{col}18", f'=IF({col}17="","",{formula_nota(col + "17")})', fmt="0.0", centro=True, bold=True)
    celda(sem, f"{col}21", "", entrada=True, wrap=True)
    sem[f"{col}21"].alignment = Alignment(wrap_text=True, vertical="top")
celda(sem, "B16", "Puntaje bruto (máx. 32)", bold=True)
celda(sem, "B17", "% de logro ponderado", bold=True)
celda(sem, "B18", "Nota del seminario", bold=True)
celda(sem, "F20", "Promedio seminarios", bold=True)
celda(sem, "G20", '=IF(COUNT(D18:G18)=0,"",ROUND(AVERAGE(D18:G18),1))', fmt="0.0", centro=True, bold=True)
sem["G20"].font = f(True, 12)
sem["H20"] = "Promedia los seminarios ya evaluados; con los 4 evaluados es la nota oficial."
sem["H20"].font = f(size=9, italic=True, color="595959")
celda(sem, "B21", "Feedback / comentarios del evaluador", bold=True)
sem.row_dimensions[21].height = 90
sem.row_dimensions[3].height = 75
sem.conditional_formatting.add("D18:G18", CellIsRule(operator="lessThan", formula=[P_NSEM], fill=relleno(ROJO)))
sem.conditional_formatting.add("D18:G18", CellIsRule(operator="greaterThanOrEqual", formula=[P_NSEM], fill=relleno(VERDE)))
sem.column_dimensions["A"].width = 5
sem.column_dimensions["B"].width = 46
sem.column_dimensions["C"].width = 8
for c in "DEFG":
    sem.column_dimensions[c].width = 24

# Desempeño
des, fin = hoja_puntaje(
    "Desempeno", "Evaluación de desempeño – Pauta de evaluación global",
    DATOS["pauta_global"], ["Mitad de rotación (formativa)", "Final (sumativa)", "Observaciones"],
    fila_titulos=[("Fecha", ["13-11-2026", "30-11-2026", ""], 4),
                  ("Evaluador", ["", "", ""], 5)],
)
# fin = 16 -> 17 suma, 18 % logro, 19 nota
for col in "DE":
    rng = f"{col}7:{col}{fin}"
    celda(des, f"{col}18", f'=IF(COUNT({rng})<{len(DATOS["pauta_global"])},"",SUMPRODUCT($C$7:$C${fin},{rng})/{P_PMAX})',
          fmt="0.0%", centro=True)
    celda(des, f"{col}19", f'=IF({col}18="","",{formula_nota(col + "18")})', fmt="0.0", centro=True, bold=True)
# la columna F es texto libre: quitar validación numérica y relleno de fórmula
des.data_validations.dataValidation[0].sqref = MultiCellRange(f"D7:E{fin}")
for rr in range(7, fin + 1):
    des[f"F{rr}"].alignment = Alignment(wrap_text=True, vertical="center")
celda(des, "B18", "% de logro ponderado", bold=True)
celda(des, "B19", "Nota (sólo la columna Final entra en la nota de la rotación)", bold=True, wrap=True)
des["E19"].font = f(True, 12)
des["D20"] = "Formativa: sin nota oficial"
des["D20"].font = f(size=9, italic=True, color="595959")
des.column_dimensions["A"].width = 5
des.column_dimensions["B"].width = 50
des.column_dimensions["C"].width = 8
des.column_dimensions["D"].width = 18
des.column_dimensions["E"].width = 18
des.column_dimensions["F"].width = 50
des.row_dimensions[6].height = 32

# ---------------------------------------------------------------- Rúbricas (descriptores)
def hoja_rubrica(nombre, tit, pauta):
    ws = wb.create_sheet(nombre)
    titulo(ws, tit, "Descriptores por nivel de desempeño. El nivel 0 corresponde a «No logrado / ausente».")
    encabezado(ws, 4, ["N°", "Criterio", "Peso", "4 · Excelente", "3 · Bueno", "2 · Suficiente", "1 · Insuficiente"])
    for i, c in enumerate(pauta):
        rr = 5 + i
        celda(ws, f"A{rr}", i + 1, centro=True)
        celda(ws, f"B{rr}", c["criterio"], bold=True, wrap=True)
        celda(ws, f"C{rr}", c["peso"], fmt="0%", centro=True)
        for j, k in enumerate(("d4", "d3", "d2", "d1")):
            celda(ws, f"{get_column_letter(4 + j)}{rr}", c[k], wrap=True)
        ws.row_dimensions[rr].height = 78
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 7
    for col in "DEFG":
        ws.column_dimensions[col].width = 38
    ws.freeze_panes = "D5"


hoja_rubrica("Rubrica_Seminario", "Rúbrica – Pauta de presentación oral (Seminarios)", DATOS["pauta_seminario"])
hoja_rubrica("Rubrica_Desempeno", "Rúbrica – Pauta de evaluación global de desempeño", DATOS["pauta_global"])

# ---------------------------------------------------------------- Bitácora
bit = wb.create_sheet("Bitacora")
titulo(bit, "Bitácora reflexiva semanal (Canvas)",
       "Preguntas: ¿Qué aprendí? · ¿Qué cambió mi forma de evaluar un paciente? · ¿Qué haría distinto ahora?")
encabezado(bit, 4, ["Semana", "Fecha límite", "Fecha de entrega", "¿Entregada a tiempo?", "Calidad reflexiva (0–4)", "Comentario del docente"])
limites = [datetime.date(2026, 11, 8), datetime.date(2026, 11, 15), datetime.date(2026, 11, 22), datetime.date(2026, 11, 29)]
for i, lim in enumerate(limites):
    rr = 5 + i
    celda(bit, f"A{rr}", f"Semana {i + 1}", centro=True)
    celda(bit, f"B{rr}", lim, fmt="dd-mm-yyyy", centro=True)
    celda(bit, f"C{rr}", None, entrada=True, fmt="dd-mm-yyyy", centro=True)
    celda(bit, f"D{rr}", f'=IF(C{rr}="","",IF(C{rr}<=B{rr},"SÍ","ATRASADA"))', centro=True)
    celda(bit, f"E{rr}", None, entrada=True, centro=True)
    celda(bit, f"F{rr}", "", entrada=True, wrap=True)
    bit.row_dimensions[rr].height = 45
dvb = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True)
bit.add_data_validation(dvb)
dvb.add("E5:E8")
celda(bit, "D10", "Entregadas", bold=True)
celda(bit, "E10", "=COUNTA(C5:C8)", centro=True)
celda(bit, "D11", "Promedio calidad", bold=True)
celda(bit, "E11", '=IF(COUNT(E5:E8)=0,"",AVERAGE(E5:E8))', fmt="0.0", centro=True)
bit["F10"] = "Referencia para puntuar el criterio 9 de la Pauta de evaluación global."
bit["F10"].font = f(size=9, italic=True, color="595959")
for col, w in zip("ABCDEF", (12, 14, 16, 20, 20, 70)):
    bit.column_dimensions[col].width = w

# ---------------------------------------------------------------- Feedback
fb = wb.create_sheet("Feedback")
titulo(fb, "Registro de feedback estructurado",
       "Modelo: 1) autoevaluación del residente · 2) fortalezas · 3) aspectos a mejorar · 4) plan de acción con compromisos")
encabezado(fb, 4, ["Fecha", "Instancia", "Evaluador", "Autoevaluación del residente", "Fortalezas observadas",
                   "Aspectos a mejorar", "Plan de acción / compromisos", "Seguimiento (¿se cumplió?)"])
instancias = [
    ("06-11-2026", "Post seminario 1"),
    ("13-11-2026", "Mitad de rotación + post seminario 2"),
    ("20-11-2026", "Post seminario 3"),
    ("27-11-2026", "Post seminario 4"),
    ("30-11-2026", "Final de rotación"),
    ("", "Otra (clínica)"),
    ("", "Otra (clínica)"),
]
for i, (fe, inst) in enumerate(instancias):
    rr = 5 + i
    celda(fb, f"A{rr}", fe, entrada=True, centro=True)
    celda(fb, f"B{rr}", inst, entrada=True, wrap=True)
    for col in "CDEFGH":
        celda(fb, f"{col}{rr}", "", entrada=True, wrap=True)
    fb.row_dimensions[rr].height = 70
for col, w in zip("ABCDEFGH", (12, 20, 16, 32, 32, 32, 32, 22)):
    fb.column_dimensions[col].width = w

# ---------------------------------------------------------------- Escala de notas
esc = wb.create_sheet("Escala_Notas")
titulo(esc, "Escala de notas (1,0 – 7,0)", "Calculada con el % de exigencia definido en «Parametros».")
encabezado(esc, 4, ["% de logro", "Nota", "", "% de logro", "Nota", "", "% de logro", "Nota"])
for k in range(101):
    bloque, pos = divmod(k, 34)
    col_p = get_column_letter(1 + 3 * bloque)
    col_n = get_column_letter(2 + 3 * bloque)
    rr = 5 + pos
    celda(esc, f"{col_p}{rr}", k / 100, fmt="0%", centro=True)
    celda(esc, f"{col_n}{rr}", "=" + formula_nota(f"{col_p}{rr}"), fmt="0.0", centro=True)
    esc.conditional_formatting.add(f"{col_n}{rr}", CellIsRule(operator="lessThan", formula=[P_APR], fill=relleno(ROJO)))
for col in "ABDEGH":
    esc.column_dimensions[col].width = 12
for col in "CF":
    esc.column_dimensions[col].width = 3

# orden de hojas
orden = ["Instrucciones", "Resumen", "Asistencia", "Seminarios", "Desempeno", "Bitacora", "Feedback",
         "Rubrica_Seminario", "Rubrica_Desempeno", "Escala_Notas", "Parametros"]
wb._sheets = [wb[n] for n in orden]
wb.active = 1
for ws in wb.worksheets:
    if ws.title in ("Resumen", "Seminarios", "Desempeno"):
        ws.sheet_properties.tabColor = AZUL_OSCURO
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

wb.save(SALIDA)
print("OK", SALIDA)
